"""Period reports (Section 11.4 P07): ``ReportEngine.compute`` is deterministic SQL over the as_of-bound views - no model.
The narrative (optional) is written by the model from the computed facts and checked by the verifier elsewhere.
Exports: PDF (reportlab), Excel (openpyxl), CSV bundle (zip), HTML - all watermarked with the synthetic-data disclaimer."""
from __future__ import annotations

import hashlib
import io
import json
import zipfile
from dataclasses import dataclass, field

import pandas as pd

from nsw_sim import CHAT_FOOTER, DISCLAIMER, clock, db
from nsw_sim.analytics import clearance, quality, queries, reconcile, statements
from nsw_sim.analytics.queries import Filters, now_iso
from nsw_sim.sim.reference import consolidated_entities

SECTIONS = ["kpis", "statements", "consolidated", "reconciliation", "clearance", "alerts", "exceptions", "origin", "scorecards", "trace_links"]
SECTION_LABELS = {"kpis": "Executive KPIs", "statements": "Per-entity statements", "consolidated": "Consolidated statement", "reconciliation": "Reconciliation summary",
                  "clearance": "Clearance performance and target tracker", "alerts": "Alerts raised and resolved", "exceptions": "Top exceptions",
                  "origin": "Origin-country breakdown", "scorecards": "Agency scorecards", "trace_links": "Appendix: trace links"}


@dataclass
class ReportParams:
    start: str
    end: str
    entities: tuple[str, ...] = ()
    filters: dict = field(default_factory=dict)
    compare: bool = True
    sections: tuple[str, ...] = tuple(SECTIONS)
    currency_view: str = "NGN"

    def to_filters(self) -> Filters:
        d = dict(self.filters)
        return Filters(start=self.start, end=self.end, entities=tuple(self.entities), currency_view=self.currency_view,
                       **{k: tuple(v) for k, v in d.items() if k in ("origins", "modes", "ports", "commodities", "processes", "fee_codes")})

    def as_dict(self) -> dict:
        return {"start": self.start, "end": self.end, "entities": list(self.entities), "filters": self.filters, "compare": self.compare,
                "sections": list(self.sections), "currency_view": self.currency_view}


def kpis(f: Filters, as_of: str | None, compare: Filters | None = None) -> dict:
    """Headline KPIs with optional comparison deltas. All values are computed by SQL/pandas."""
    def one(ff: Filters) -> dict:
        mix = queries.entity_mix(ff, as_of)
        a, p, s = mix["assessed"].sum(), mix["paid"].sum(), mix["settled"].sum()
        cc = mix["collection_cost"].sum()
        dw = clearance.dwell_table(ff, None, as_of)
        return {"assessed": float(a), "paid": float(p), "settled": float(s), "remitted": float(mix["remitted"].sum()), "outstanding": float(mix["outstanding"].sum()),
                "in_transit": float(mix["in_transit"].sum()), "collection_cost": float(cc), "cost_per_100": float(100 * cc / (s + cc)) if (s + cc) else None,
                "median_dwell_days": (float(dw["dwell_p50"].iloc[0]) if not dw.empty and dw["dwell_p50"].iloc[0] == dw["dwell_p50"].iloc[0] else None),
                "sla_breach_rate": clearance.overall_sla_breach(ff, as_of),
                "consignments": int(queries.total("consignments", ff, as_of)),
                "open_high_alerts": int(db.reader(as_of).execute("SELECT COUNT(*) FROM v_alerts WHERE severity='high' AND status IN ('open','acknowledged','under_review')").fetchone()[0]),
                "confidence_score": quality.overall_score(as_of)}
    cur = one(f)
    prev = one(compare) if compare else None
    return {"current": cur, "previous": prev}


class ReportEngine:
    @staticmethod
    def compute(params: ReportParams, as_of: str | None = None) -> dict:
        as_of = as_of or now_iso()
        f = params.to_filters()
        f = f.with_(end=min(f.end, as_of))
        prev = queries.previous_period(f) if params.compare else None
        ents = list(params.entities) or consolidated_entities()
        out: dict = {"params": params.as_dict(), "as_of": as_of, "period": [f.start, f.end], "disclaimer": DISCLAIMER}
        S = set(params.sections)
        if "kpis" in S:
            out["kpis"] = kpis(f, as_of, prev)
        if "statements" in S:
            out["statements"] = {}
            for e in ents:
                out["statements"][e] = {"performance": statements.performance(e, f, as_of, prev),
                                        "position": statements.position(e, f.end, as_of, prev.end if prev else None),
                                        "cash_flow": statements.cash_flow(e, f, as_of),
                                        "collection": statements.collection_remittance(e, f, as_of)}
        if "consolidated" in S:
            out["consolidated"] = {"performance": statements.performance("ALL", f, as_of, prev), "position": statements.position("ALL", f.end, as_of, prev.end if prev else None),
                                   "cash_flow": statements.cash_flow("ALL", f, as_of), "collection": statements.collection_remittance("ALL", f, as_of)}
        if "reconciliation" in S:
            out["reconciliation"] = {"four_way": reconcile.four_way(f, as_of), "summary": reconcile.exception_summary(f, as_of), "bridge": reconcile.variance_bridge(f, as_of),
                                     "heatmap": reconcile.leakage_heatmap(f, as_of).head(15), "ageing": reconcile.unpaid_ageing(f, as_of)}
        if "clearance" in S:
            wk = clearance.weekly_dwell(f.with_(start=None), as_of)
            from nsw_sim.analytics import forecast
            out["clearance"] = {"dwell": clearance.dwell_table(f, None, as_of), "stages": clearance.stage_summary(f, as_of),
                                "bottlenecks": clearance.bottlenecks(f, as_of), "digital_share": clearance.overall_digital_share(f, as_of),
                                "weekly": wk, "projection": {k: v for k, v in forecast.project_to_target(wk).items() if k != "series"}}
        if "alerts" in S:
            out["alerts"] = pd.read_sql_query("SELECT alert_id, rule_code, severity, entity_id, subject, detected_at, status FROM v_alerts WHERE detected_at>=? AND detected_at<? ORDER BY detected_at",
                                              db.reader(as_of), params=(f.start, f.end))
        if "exceptions" in S:
            out["exceptions"] = reconcile.exceptions(f, as_of, 25).head(25)
        if "origin" in S:
            out["origin"] = queries.country_collections(f, "assessed", as_of).head(15)
        if "scorecards" in S:
            out["scorecards"] = quality.scorecard(as_of)
        if "trace_links" in S:
            ex = out.get("exceptions")
            out["trace_links"] = [] if ex is None or ex.empty else [{"ref": r.nsw_ref, "class": r.cls, "entity": r.entity} for r in ex.head(10).itertuples() if r.nsw_ref]
        out["fingerprint"] = fingerprint(out)
        return out


def fingerprint(report: dict) -> str:
    """SHA-256 of (parameters, as_of, row counts, totals): identical for identical inputs at the same as_of."""
    basis = {"params": report["params"], "as_of": report["as_of"], "totals": {}}
    t = basis["totals"]
    if "kpis" in report:
        t["kpis"] = {k: (None if v is None else round(float(v), 6)) for k, v in report["kpis"]["current"].items() if k != "confidence_score"}
    if "consolidated" in report:
        perf = report["consolidated"]["performance"]
        t["consolidated"] = {"revenue": perf["total_revenue"], "expenses": perf["total_expenses"], "surplus": perf["surplus"],
                             "assets": report["consolidated"]["position"]["total_assets"]}
    if "statements" in report:
        t["statements"] = {e: [s["performance"]["total_revenue"], s["performance"]["surplus"], s["position"]["total_assets"], s["cash_flow"]["closing"]] for e, s in report["statements"].items()}
    if "reconciliation" in report:
        t["recon"] = report["reconciliation"]["four_way"]
        t["recon_rows"] = int(len(report["reconciliation"]["summary"]))
    for key in ("alerts", "exceptions", "origin", "scorecards"):
        if key in report:
            t[key + "_rows"] = int(len(report[key]))
    if "clearance" in report:
        t["clearance_rows"] = [int(len(report["clearance"]["stages"])), int(len(report["clearance"]["weekly"]))]
    return hashlib.sha256(json.dumps(basis, sort_keys=True, default=str).encode()).hexdigest()


# ------------------------------------------------------------------------------- exports
def _flat_tables(rep: dict) -> dict[str, pd.DataFrame]:
    """Every section as a DataFrame (for Excel sheets and the CSV bundle)."""
    t: dict[str, pd.DataFrame] = {}
    if "kpis" in rep:
        cur, prev = rep["kpis"]["current"], rep["kpis"]["previous"] or {}
        t["kpis"] = pd.DataFrame([{"kpi": k, "value": v, "previous_period": prev.get(k)} for k, v in cur.items()])
    for _scope, key in (("consolidated", "consolidated"),):
        if key in rep:
            t["consolidated_performance"] = _lines_df(rep[key]["performance"]["revenue"] + rep[key]["performance"]["expenses"])
            t["consolidated_position"] = _lines_df(rep[key]["position"]["assets"] + rep[key]["position"]["liabilities"] + rep[key]["position"]["equity"])
            t["consolidated_cash_flow"] = pd.DataFrame(rep[key]["cash_flow"]["lines"])[["label", "amount"]]
            t["consolidated_collection"] = rep[key]["collection"]
    for ent, s in rep.get("statements", {}).items():
        t[f"{ent}_performance"] = _lines_df(s["performance"]["revenue"] + s["performance"]["expenses"])
        t[f"{ent}_position"] = _lines_df(s["position"]["assets"] + s["position"]["liabilities"] + s["position"]["equity"])
    if "reconciliation" in rep:
        r = rep["reconciliation"]
        t["recon_funnel"] = pd.DataFrame([{"stage": k, **v} for k, v in r["four_way"].items()])
        t["recon_exceptions"] = r["summary"]
        t["recon_bridge"] = pd.DataFrame(r["bridge"])
        t["leakage_heatmap"] = r["heatmap"]
    if "clearance" in rep:
        c = rep["clearance"]
        t["dwell_percentiles"], t["stage_summary"], t["bottlenecks"], t["weekly_dwell"] = c["dwell"], c["stages"], c["bottlenecks"], c["weekly"]
    for k in ("alerts", "exceptions", "origin", "scorecards"):
        if k in rep:
            t[k] = rep[k]
    return {k: v for k, v in t.items() if isinstance(v, pd.DataFrame)}


def _lines_df(lines: list[dict]) -> pd.DataFrame:
    return pd.DataFrame([{"section": x["section"], "line": x["label"], "amount_ngn": x["amount"] / 100.0, "previous_period_ngn": None if x.get("prev") is None else x["prev"] / 100.0} for x in lines])


def to_excel(rep: dict) -> bytes:
    """One sheet per section plus a metadata sheet repeating the disclaimer; the disclaimer also heads every sheet."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Metadata"
    for row in (["NSW Intelligence Console - period report"], [DISCLAIMER], [], ["Period (UTC)", rep["period"][0], rep["period"][1]], ["As of", rep["as_of"]],
                ["Report fingerprint (SHA-256)", rep["fingerprint"]], ["Currency view", rep["params"]["currency_view"]], ["Amounts", "Naira (₦) unless the column says otherwise"]):
        ws.append(row)
    ws["A2"].font = Font(bold=True, color="B03A2E")
    for name, df in _flat_tables(rep).items():
        sh = wb.create_sheet(name[:31])
        sh.append([DISCLAIMER])
        sh["A1"].font = Font(bold=True, color="B03A2E")
        sh.append([])
        sh.append(list(df.columns))
        for r in df.itertuples(index=False):
            sh.append([None if (isinstance(v, float) and v != v) else (v if isinstance(v, (int, float, str)) or v is None else str(v)) for v in r])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_csv_zip(rep: dict) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("README_DISCLAIMER.txt", f"{DISCLAIMER}\nFingerprint: {rep['fingerprint']}\nAs of: {rep['as_of']}\n")
        for name, df in _flat_tables(rep).items():
            z.writestr(f"{name}.csv", f"# {DISCLAIMER}\n" + df.to_csv(index=False))
    return buf.getvalue()


def to_html(rep: dict, narrative: str | None = None) -> str:
    body = [f"<h1>NSW Intelligence Console: period report</h1><p style='color:#B03A2E;font-weight:700'>{DISCLAIMER}</p>",
            f"<p>Period (UTC): {rep['period'][0]} → {rep['period'][1]} · as of {rep['as_of']}<br>Fingerprint: <code>{rep['fingerprint']}</code></p>"]
    if narrative:
        body.append(f"<h2>Executive summary</h2><p>{narrative}</p>")
    for name, df in _flat_tables(rep).items():
        body.append(f"<h3>{name.replace('_', ' ')}</h3>" + df.head(200).to_html(index=False, border=0, float_format=lambda x: f"{x:,.2f}"))
    css = "body{font-family:Arial,sans-serif;margin:24px;color:#1B2420}table{border-collapse:collapse;font-size:12px}td,th{border-bottom:1px solid #ddd;padding:3px 8px;text-align:right}" \
          "th{background:#0B5D3B;color:#fff}.wm{position:fixed;top:40%;left:5%;font-size:64px;color:rgba(176,58,46,.08);transform:rotate(-20deg)}"
    return f"<html><head><meta charset='utf-8'><title>NSW period report</title><style>{css}</style></head><body><div class='wm'>SYNTHETIC DATA</div>{''.join(body)}<p>{CHAT_FOOTER}</p></body></html>"


def to_pdf(rep: dict, narrative: str | None = None) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    def deco(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(colors.Color(0.69, 0.23, 0.18, alpha=0.10))
        canvas.setFont("Helvetica-Bold", 60)
        canvas.translate(420, 300)
        canvas.rotate(25)
        canvas.drawCentredString(0, 0, "SYNTHETIC DATA")
        canvas.restoreState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#B03A2E"))
        canvas.drawString(30, 18, f"{DISCLAIMER} · fingerprint {rep['fingerprint'][:16]}… · page {doc.page}")
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=30, rightMargin=30, topMargin=30, bottomMargin=34)
    st = getSampleStyleSheet()
    els = [Paragraph("NSW Intelligence Console: period report", st["Title"]), Paragraph(DISCLAIMER, st["Heading3"]),
           Paragraph(f"Period (UTC) {rep['period'][0]} to {rep['period'][1]} · as of {rep['as_of']}<br/>Fingerprint (SHA-256): {rep['fingerprint']}", st["Normal"]), Spacer(1, 10)]
    if narrative:
        els += [Paragraph("Executive summary", st["Heading2"]), Paragraph(narrative.replace("\n", "<br/>"), st["Normal"]), Spacer(1, 8)]
    for name, df in _flat_tables(rep).items():
        els.append(Paragraph(name.replace("_", " ").title(), st["Heading3"]))
        d = df.head(30)
        data = [list(map(str, d.columns))] + [[(f"{v:,.2f}" if isinstance(v, float) else str(v))[:34] for v in row] for row in d.itertuples(index=False)]
        t = Table(data, repeatRows=1)
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0B5D3B")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTSIZE", (0, 0), (-1, -1), 6.5),
                               ("GRID", (0, 0), (-1, -1), 0.25, colors.lightgrey), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F5F3")])]))
        els += [t, Spacer(1, 8)]
    els.append(Paragraph(CHAT_FOOTER, st["Italic"]))
    doc.build(els, onFirstPage=deco, onLaterPages=deco)
    return buf.getvalue()


def save_report(conn, rep: dict, creator: str, files: dict[str, str]) -> str:
    rid = f"RPT-{clock.wat_day(rep['as_of']).replace('-', '')}-{rep['fingerprint'][:6].upper()}"
    conn.execute("INSERT OR REPLACE INTO saved_reports(report_id,created_at,creator,params_json,as_of,fingerprint,status,file_paths_json) VALUES(?,?,?,?,?,?,?,?)",
                 (rid, now_iso(), creator, json.dumps(rep["params"]), rep["as_of"], rep["fingerprint"], "awaiting sign-off", json.dumps(files)))
    return rid
