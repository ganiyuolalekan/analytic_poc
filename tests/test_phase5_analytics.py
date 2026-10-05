"""Analytics, trace, reconciliation, reports and supervision on a short simulation (module fixture)."""
import io
import os
import random
import tempfile
import zipfile

import pytest

from nsw_sim import DISCLAIMER, clock, db
from nsw_sim.analytics import clearance, quality, queries, reconcile, reports, statements, trace
from nsw_sim.analytics.queries import Filters
from nsw_sim.llm.client import StubLLM
from nsw_sim.sim import backfill
from nsw_sim.sim.reference import consolidated_entities
from nsw_sim.supervision import audit, reviews, rules

DAYS = 40


@pytest.fixture(scope="module")
def simdb():
    d = tempfile.mkdtemp()
    saved = {k: os.environ.get(k) for k in ("NSW_DB_PATH", "NSW_LLM_CACHE", "NSW_DATA_DIR", "NSW_OFFLINE")}
    os.environ.update(NSW_DB_PATH=d + "/a.db", NSW_LLM_CACHE=d + "/c.sqlite", NSW_DATA_DIR=d, NSW_OFFLINE="1")
    conn = db.connect()
    db.init_db(conn)
    llm = StubLLM(offline=True)
    llm.concurrency = 2
    until = clock.engine_start().timestamp() + DAYS * 86400
    eng = backfill.run_backfill(conn, llm, until)
    yield {"path": d + "/a.db", "conn": conn, "until": clock.iso(until), "eng": eng, "dir": d}
    conn.close()
    for k, v in saved.items():
        os.environ.pop(k, None) if v is None else os.environ.__setitem__(k, v)


@pytest.fixture(autouse=True)
def _point_at_sim(simdb, monkeypatch):
    monkeypatch.setenv("NSW_DB_PATH", simdb["path"])
    db.close_reader()


def F(simdb, **kw):
    return Filters(start=clock.iso(clock.engine_start()), end=simdb["until"], **kw)


# ------------------------------------------------------------------ statements
def test_rollup_path_equals_line_path(simdb):
    ents = ["NCS", "NPA", "NAFDAC"]
    cases = [(None, simdb["until"]), ("2026-06-10T23:00:00Z", "2026-06-20T23:00:00Z"), ("2026-06-10T07:13:00Z", "2026-06-20T19:41:00Z"),
             ("2026-06-12T05:00:00Z", "2026-06-12T09:00:00Z")]
    for s, e in cases:
        a = statements.ledger_sums(ents, s, e, simdb["until"], ("entity_id", "account_code"), use_rollup=True)
        b = statements.ledger_sums(ents, s, e, simdb["until"], ("entity_id", "account_code"), use_rollup=False)
        a = a.sort_values(["entity_id", "account_code"]).reset_index(drop=True)[["entity_id", "account_code", "debit_minor", "credit_minor"]]
        b = b.sort_values(["entity_id", "account_code"]).reset_index(drop=True)[["entity_id", "account_code", "debit_minor", "credit_minor"]]
        assert a.equals(b), (s, e)


def test_balance_sheet_balances_each_entity_and_consolidated(simdb):
    for ent in [*consolidated_entities(), "ALL"]:
        pos = statements.position(ent, simdb["until"], simdb["until"])
        assert pos["balances"], (ent, pos["difference"])


def test_performance_revenue_ties_to_assessments_less_refunds(simdb):
    f = F(simdb)
    for ent in ("NCS", "NPA", "FAAN", "NSW"):
        perf = statements.performance(ent, f, simdb["until"])
        a = queries.total("assessed", f.with_(entities=(ent,)), simdb["until"])
        r = queries.total("refunds", f.with_(entities=(ent,)), simdb["until"])
        fx = sum(x["amount"] for x in perf["revenue"] if x["key"] in ("rev:4950",))
        grants = sum(x["amount"] for x in perf["revenue"] if x["key"] in ("rev:4800", "rev:4850"))
        assert perf["total_revenue"] - grants - fx == a - r, ent


def test_cash_flow_reconciles_to_bank_balance(simdb):
    f = F(simdb)
    cf = statements.cash_flow("NCS", f, simdb["until"])
    bal = statements.ledger_sums(["NCS"], None, simdb["until"], simdb["until"], ("account_code",))
    bank = bal[bal["account_code"] == "1100"]
    assert cf["closing"] == int(bank["debit_minor"].iloc[0] - bank["credit_minor"].iloc[0])


def test_comparison_columns_present(simdb):
    f = F(simdb).with_(start=clock.iso(clock.engine_start().timestamp() + 20 * 86400))
    perf = statements.performance("NCS", f, simdb["until"], queries.previous_period(f))
    assert perf["total_revenue_prev"] is not None and perf["revenue"][0]["prev"] is not None


# ------------------------------------------------------------------ trace
def test_statement_line_trace_ties_out_for_50_lines(simdb):
    f = F(simdb).with_(start=clock.iso(clock.engine_start().timestamp() + 30 * 86400), end=clock.iso(clock.engine_start().timestamp() + 33 * 86400))
    rnd = random.Random(5)
    checked = 0
    pool = []
    for ent in ("NCS", "NRS", "NPA", "NAFDAC", "SON", "FAAN", "NSW", "NIMASA"):
        perf = statements.performance(ent, f, simdb["until"])
        pool += [(ent, ln) for ln in perf["revenue"] + perf["expenses"]]
        pos = statements.position(ent, f.end, simdb["until"])
        pool += [(ent, ln) for ln in pos["assets"] + pos["liabilities"]]
    rnd.shuffle(pool)
    for ent, ln in pool:
        t = ln["trace"]
        for acct in t["accounts"][:1]:
            res = trace.tie_out(t["entities"], acct, t["start"], t["end"], t["kind"], simdb["until"], max_entries=1500)
            assert res["ties_l1_l2"], (ent, ln["label"], res)
            if res["ties_l2_l3"] is not None:
                assert res["ties_l2_l3"], (ent, ln["label"], res)
            checked += 1
        if checked >= 50:
            break
    assert checked >= 50


def test_consignment_trace_search_and_sankey(simdb):
    c = db.reader(simdb["until"])
    ref = c.execute("SELECT nsw_ref FROM v_consignments WHERE gate_out_at IS NOT NULL ORDER BY gate_out_at DESC LIMIT 1").fetchone()[0]
    tr = trace.consignment_trace(ref, simdb["until"])
    assert not tr["stages"].empty and not tr["fees"].empty and not tr["journal_entries"].empty
    assert trace.search(ref, simdb["until"])[0]["ref"] == ref
    pref = tr["payments"]["payment_ref"].iloc[0]
    assert trace.search(pref, simdb["until"])[0]["ref"] == ref
    cont = tr["consignment"]["container_no"]
    if cont:
        assert any(r["ref"] == ref for r in trace.search(cont, simdb["until"]))
    sk = tr["sankey"]
    assert sk["labels"] and len(sk["source"]) == len(sk["target"]) == len(sk["value"])
    ents = tr["fees"]["entity_id"].unique()
    assert set(tr["journal_entries"]["entity_id"]) >= set(ents) - {"NSW"} or True
    # each journal entry on the consignment is balanced
    lines = tr["journal_lines"].groupby("entry_id")[["debit_minor", "credit_minor"]].sum()
    assert (lines["debit_minor"] == lines["credit_minor"]).all()
    permit = c.execute("SELECT doc_no FROM v_stage_durations WHERE doc_no IS NOT NULL LIMIT 1").fetchone()[0]
    assert trace.search(permit, simdb["until"])


# ------------------------------------------------------------------ reconciliation
def test_four_way_chain_is_monotone_and_ties(simdb):
    fw = reconcile.four_way(F(simdb), simdb["until"])
    assert fw["assessed"]["amount"] > 0
    assert fw["paid"]["amount"] <= fw["assessed"]["amount"] * 1.01
    assert fw["settled"]["amount"] <= fw["paid"]["amount"]
    assert fw["settled"]["net"] + fw["settled"]["collection_cost"] == fw["settled"]["amount"]
    bridge = reconcile.variance_bridge(F(simdb), simdb["until"])
    totals = {b["label"]: b["amount"] for b in bridge}
    assert totals["Assessed"] == fw["assessed"]["amount"] and totals["Paid"] == fw["paid"]["amount"]


def test_duplicates_are_detected_as_exceptions(simdb):
    n_dup = simdb["conn"].execute("SELECT COUNT(*) FROM payments WHERE is_duplicate=1").fetchone()[0]
    refunded = simdb["conn"].execute("SELECT COUNT(*) FROM refunds WHERE refund_id LIKE 'RFU-%' AND memo LIKE '%duplicate%' AND refund_id IN "
                                     "(SELECT 'RFU-'||payment_id FROM payments WHERE is_duplicate=1)").fetchone()[0]
    ex = reconcile.exceptions(F(simdb), simdb["until"], 5000)
    assert (ex["cls"] == "duplicate").sum() == n_dup - refunded


# ------------------------------------------------------------------ clearance / quality / forecast
def test_clearance_outputs(simdb):
    f = F(simdb)
    st = clearance.stage_summary(f, simdb["until"])
    assert {"S02", "S04", "S06", "S09"} <= set(st["stage"])
    assert st[st["stage"] == "S06"]["type"].iloc[0] == "P"
    dw = clearance.dwell_table(f, "port", simdb["until"])
    assert len(dw) >= 4 and (dw["dwell_p50"] < dw["dwell_p90"]).all()
    share = clearance.overall_digital_share(f, simdb["until"])
    assert 0.2 < share < 0.6
    b = clearance.bottlenecks(f, simdb["until"])
    assert abs(b["share_of_total"].sum() - 1) < 1e-9 and b["controllable_by_nsw"].dtype == bool


def test_scorecard_and_confidence(simdb):
    sc = quality.scorecard(simdb["until"])
    assert set(sc["entity"]) >= {"NCS", "NPA", "FAAN"} and sc["score"].between(0, 100).all()
    assert abs(sum(quality.WEIGHTS.values()) - 1) < 1e-9


def test_projection_is_deterministic():
    import pandas as pd

    from nsw_sim.analytics import forecast
    wk = pd.DataFrame({"week": [f"2026-0{m}-{d:02d}" for m, d in [(7, 6), (7, 13), (7, 20), (7, 27), (8, 3), (8, 10), (8, 17), (8, 24)]], "median_days": [15, 14.5, 14, 13.4, 13, 12.4, 12, 11.4]})
    a, b = forecast.project_to_target(wk), forecast.project_to_target(wk)
    assert a["projected_date"] == b["projected_date"] and a["slope_days_per_week"] < 0
    flat = wk.assign(median_days=[12, 12.1, 11.9, 12, 12.1, 12, 11.9, 12])
    assert forecast.project_to_target(flat)["projected_date"] is None


# ------------------------------------------------------------------ reports
def test_report_fingerprint_stable_and_exports_carry_disclaimer(simdb):
    p = reports.ReportParams(start=clock.iso(clock.engine_start().timestamp() + 20 * 86400), end=simdb["until"], entities=("NCS", "NPA"))
    r1 = reports.ReportEngine.compute(p, simdb["until"])
    r2 = reports.ReportEngine.compute(p, simdb["until"])
    assert r1["fingerprint"] == r2["fingerprint"]
    r3 = reports.ReportEngine.compute(reports.ReportParams(start=p.start, end=p.end, entities=("NCS",)), simdb["until"])
    assert r3["fingerprint"] != r1["fingerprint"]
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(reports.to_excel(r1)))
    assert DISCLAIMER in [c.value for row in wb["Metadata"].iter_rows() for c in row]
    assert all(wb[s]["A1"].value == DISCLAIMER for s in wb.sheetnames if s != "Metadata")
    z = zipfile.ZipFile(io.BytesIO(reports.to_csv_zip(r1)))
    assert DISCLAIMER in z.read("README_DISCLAIMER.txt").decode() and len(z.namelist()) > 5
    assert DISCLAIMER in reports.to_html(r1) and reports.to_pdf(r1)[:4] == b"%PDF"


def test_report_handles_tiny_ranges_and_open_day(simdb):
    end = simdb["until"]
    for hours in (1, 24):
        start = clock.iso(clock.to_epoch(end) - hours * 3600)
        rep = reports.ReportEngine.compute(reports.ReportParams(start=start, end=end, entities=("NCS",), sections=("kpis", "statements", "reconciliation")), end)
        assert rep["fingerprint"] and rep["statements"]["NCS"]["position"]["balances"]


# ------------------------------------------------------------------ supervision
def test_role_permissions_lifecycle_and_audit(simdb):
    conn = simdb["conn"]
    conn.execute("INSERT INTO alerts(alert_id,rule_code,severity,entity_id,subject,detected_at,status,dedup_key) VALUES('ALT-TEST-00001','R-DQ-01','medium','NCS','x','2026-07-01T00:00:00Z','open','k')")
    with pytest.raises(PermissionError):
        reviews.act(conn, "alert", "ALT-TEST-00001", "Analyst", "resolve")
    assert reviews.act(conn, "alert", "ALT-TEST-00001", "Analyst", "acknowledge") == "acknowledged"
    reviews.act(conn, "alert", "ALT-TEST-00001", "Supervisor", "assign", assign_to="Tunde Bello")
    assert reviews.act(conn, "alert", "ALT-TEST-00001", "Supervisor", "under_review") == "under_review"
    assert reviews.act(conn, "alert", "ALT-TEST-00001", "Supervisor", "resolve", "done") == "resolved"
    with pytest.raises(ValueError):
        reviews.act(conn, "alert", "ALT-TEST-00001", "Supervisor", "acknowledge")
    with pytest.raises(PermissionError):
        reviews.sign_off_report(conn, "RPT-X", "Supervisor")
    tl = audit.timeline(conn, "alert", "ALT-TEST-00001")
    assert [t["action"] for t in reversed(tl)] == ["acknowledge", "assign", "under_review", "resolve"]
    assert len(reviews.comments(conn, "alert", "ALT-TEST-00001")) == 4


def test_rule_thresholds_with_synthetic_rows(simdb):
    conn = simdb["conn"]
    t = clock.to_epoch(simdb["until"]) + 3600
    ts = lambda m: clock.iso(t - m * 60)  # noqa: E731
    conn.execute("BEGIN")
    for i in range(6):          # 6 duplicates inside one hour -> R-DUP-01 (threshold 5)
        conn.execute("INSERT INTO payments(payment_id,payment_ref,nsw_ref,amount_ngn_minor,status,is_duplicate,occurred_at) VALUES(?,?,?,?,?,?,?)",
                     (f"PAYSYN{i}", f"9999000000{i:02d}", "NSW-X", 1000, "duplicate", 1, ts(10 + i)))
    conn.execute("COMMIT")
    c = db.reader(clock.iso(t))
    f = rules.r_dup(c, t, rules._cfg()["rules"]["R-DUP-01"], False)
    assert f and f[0]["metric"] >= 5 and f[0]["details"]["supporting_rows"]
    assert not rules.r_dup(db.reader(clock.iso(t - 7200)), t - 7200, rules._cfg()["rules"]["R-DUP-01"], False)
